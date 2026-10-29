import logfire
from docx import Document
from pptx import Presentation

def _docx_lines(file_path: str):
    doc = Document(file_path)
    for p in doc.paragraphs:
        yield p.text
    for table in doc.tables:
        for row in table.rows:
            yield " | ".join(cell.text for cell in row.cells)

def _pptx_lines(file_path: str):
    prs = Presentation(file_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    yield "".join(run.text for run in p.runs)
            if shape.has_table:
                for row in shape.table.rows:
                    yield " | ".join(cell.text for cell in row.cells)

def parse_office(file_path: str):
    """
    Parses Office documents (.docx, .pptx) with python-docx / python-pptx.
    Unlike PDFs, these formats are structured and lightweight, so they are processed locally.
    (Avoids Unstructured's libmagic dependency, which crashes on Windows.)
    """
    with logfire.span("📄 Office Document Parsing", filename=file_path):
        try:
            ext = file_path.lower().rsplit('.', 1)[-1]
            lines = _pptx_lines(file_path) if ext == 'pptx' else _docx_lines(file_path)
            full_text = "\n".join(line for line in lines if line.strip())

            if not full_text.strip():
                logfire.warning(f"⚠️ No text extracted from {file_path}")
            else:
                logfire.info(f"✅ Successfully parsed {len(full_text)} characters")

            return full_text
        except Exception as e:
            logfire.error(f"❌ Office Parse Failed: {e}")
            raise e
